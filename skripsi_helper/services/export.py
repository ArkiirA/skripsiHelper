"""Export service: generate DOCX and PDF from skripsi project."""
from pathlib import Path
from typing import Optional, List, Any
import subprocess
import tempfile
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE

from skripsi_helper.models.skripsi import SkripsiProject, Section
from skripsi_helper.models.pedoman import PedomanPreset
from skripsi_helper.models.citation import Citation
from skripsi_helper.services.citation_extractor import format_citation_string


def build_docx_document(
    project: Any,
    sections: List[Any],
    citations: List[Any],
    pedoman: Optional[Any] = None,
    output_path: Optional[Path] = None,
) -> bytes:
    """Build DOCX document from skripsi content."""
    doc = Document()

    # Formatting rules from pedoman or defaults
    margins = {"top": 3.0, "bottom": 3.0, "left": 4.0, "right": 3.0}
    font_name = "Times New Roman"
    font_size = 12
    line_spacing = 1.5

    if pedoman is not None:
        if getattr(pedoman, "default_margins", None):
            margins = pedoman.default_margins
        if getattr(pedoman, "default_font", None):
            font_name = pedoman.default_font
        if getattr(pedoman, "default_font_size", None):
            font_size = pedoman.default_font_size
        if getattr(pedoman, "default_line_spacing", None):
            line_spacing = pedoman.default_line_spacing

    # Set document margins (1 cm = 0.3937 inches)
    doc_sections = doc.sections
    doc_section = doc_sections[0]
    doc_section.top_margin = Inches(margins.get("top", 3.0) * 0.3937)
    doc_section.bottom_margin = Inches(margins.get("bottom", 3.0) * 0.3937)
    doc_section.left_margin = Inches(margins.get("left", 4.0) * 0.3937)
    doc_section.right_margin = Inches(margins.get("right", 3.0) * 0.3937)

    # Title page
    title_text = str(getattr(project, "title", "SKRIPSI"))
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title_text.upper())
    run.font.name = font_name
    run.font.size = Pt(font_size + 2)
    run.bold = True

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = sub_para.add_run("\n\nSKRIPSI\n\n\n\n")
    run2.font.name = font_name
    run2.font.size = Pt(font_size)

    # Page break after title page
    doc.add_page_break()

    # Table of contents placeholder
    toc_heading = doc.add_paragraph()
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_toc = toc_heading.add_run("DAFTAR ISI\n\n")
    run_toc.font.name = font_name
    run_toc.font.size = Pt(font_size + 2)
    run_toc.bold = True

    # Page break after TOC
    doc.add_page_break()

    # Sections sorted by order
    sorted_sections = sorted(sections, key=lambda s: getattr(s, "order", 0))
    for sec in sorted_sections:
        sec_title = str(getattr(sec, "title", ""))
        sec_heading = doc.add_heading(sec_title, level=1)
        sec_heading.paragraph_format.space_before = Pt(12)
        sec_heading.paragraph_format.space_after = Pt(6)

        # Parse content
        content_val = getattr(sec, "content", None)
        content_dict = content_val if isinstance(content_val, dict) else {}
        content_list = content_dict.get("content", []) if content_dict else []

        for block in content_list:
            block_type = block.get("type", "paragraph")
            text_content = ""

            if block_type == "paragraph":
                for inline in block.get("content", []):
                    if inline.get("type") == "text":
                        text_content += inline.get("text", "")
                if text_content:
                    para = doc.add_paragraph(text_content)
                    para_format = para.paragraph_format
                    para_format.line_spacing = line_spacing
                    para_format.space_after = Pt(6)

            elif block_type in ["heading1", "heading2", "heading3"]:
                level = 1 if block_type == "heading1" else 2 if block_type == "heading2" else 3
                for inline in block.get("content", []):
                    if inline.get("type") == "text":
                        text_content += inline.get("text", "")
                if text_content:
                    doc.add_heading(text_content, level=level)

    # Bibliography
    if citations:
        doc.add_page_break()
        bib_heading = doc.add_paragraph()
        bib_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_bib = bib_heading.add_run("DAFTAR PUSTAKA\n\n")
        run_bib.font.name = font_name
        run_bib.font.size = Pt(font_size + 2)
        run_bib.bold = True

        for i, citation in enumerate(citations, 1):
            cite_text = getattr(citation, "citation_string", None) or ""
            if not cite_text:
                cite_data = {
                    "authors": getattr(citation, "authors", []),
                    "title": getattr(citation, "title", ""),
                    "year": getattr(citation, "year", None),
                    "journal": getattr(citation, "journal", None),
                    "volume": getattr(citation, "volume", None),
                    "issue": getattr(citation, "issue", None),
                    "pages": getattr(citation, "pages", None),
                    "doi": getattr(citation, "doi", None),
                }
                cite_text = format_citation_string(cite_data)

            para = doc.add_paragraph(f"[{i}] {cite_text}")
            para_format = para.paragraph_format
            para_format.left_indent = Inches(0.5)
            para_format.line_spacing = line_spacing
            para_format.space_after = Pt(3)

    from io import BytesIO
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    if output_path:
        with open(output_path, "wb") as f:
            f.write(buffer.getvalue())

    return buffer.getvalue()


def convert_docx_to_pdf(docx_bytes: bytes) -> bytes:
    """Convert DOCX bytes to PDF using LibreOffice headless."""
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = Path(tmpdir) / "document.docx"
        pdf_path = Path(tmpdir) / "document.pdf"

        with open(docx_path, "wb") as f:
            f.write(docx_bytes)

        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to", "pdf",
            "--outdir", str(tmpdir),
            str(docx_path),
        ]

        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)

        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")

        if not pdf_path.exists():
            raise RuntimeError(f"PDF file was not created: {result.stderr}")

        with open(pdf_path, "rb") as f:
            return f.read()
